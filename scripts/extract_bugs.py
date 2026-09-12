#!/usr/bin/env python3
"""Extract all bugs from the test execution tracker Excel."""
import openpyxl

wb = openpyxl.load_workbook('/home/z/my-project/upload/3BOXES HRMS_Test_Execution_Tracker (1).xlsx', data_only=True)

output_lines = []

for sheet_name in ['Employee Management', 'Leave Management', 'Attendance Module']:
    ws = wb[sheet_name]
    output_lines.append(f'\n{"="*80}')
    output_lines.append(f'SHEET: {sheet_name}')
    output_lines.append(f'{"="*80}')

    # Two tracker sections: cols 1-9 (left) and cols 12-20 (right)
    for section_start, section_label in [(1, 'LEFT (3boxeshrms.com)'), (12, 'RIGHT (marqaitechgroup)')]:
        output_lines.append(f'\n--- Section: {section_label} ---')

        for row in range(3, ws.max_row + 1):
            test_id = ws.cell(row=row, column=section_start).value
            if not test_id or str(test_id).strip() == '' or str(test_id).strip() == 'None':
                continue
            module = ws.cell(row=row, column=section_start + 1).value or ''
            feature = ws.cell(row=row, column=section_start + 2).value or ''
            scenario = ws.cell(row=row, column=section_start + 3).value or ''
            steps = ws.cell(row=row, column=section_start + 4).value or ''
            test_data = ws.cell(row=row, column=section_start + 5).value or ''
            expected = ws.cell(row=row, column=section_start + 6).value or ''
            actual = ws.cell(row=row, column=section_start + 7).value or ''
            status = ws.cell(row=row, column=section_start + 8).value or ''

            # Only capture FAIL / PENDING / issues
            status_str = str(status).strip().upper()
            actual_str = str(actual).strip()
            if any(kw in status_str for kw in ['FAIL', 'PENDING', 'BLOCK', 'BUG', 'ISSUE', 'NOT PASS', 'INCOMPLETE', 'RETEST']):
                output_lines.append(f'\n  Row {row} | TestID: {test_id} | Status: {status}')
                output_lines.append(f'    Module: {module}')
                output_lines.append(f'    Feature: {feature}')
                output_lines.append(f'    Scenario: {scenario}')
                output_lines.append(f'    Steps: {steps}')
                output_lines.append(f'    Test Data: {test_data}')
                output_lines.append(f'    Expected: {expected}')
                output_lines.append(f'    Actual: {actual}')

# Write to file
with open('/home/z/my-project/scripts/extracted_bugs.txt', 'w') as f:
    f.write('\n'.join(output_lines))

print('Bugs extracted to /home/z/my-project/scripts/extracted_bugs.txt')
print(f'Total lines: {len(output_lines)}')
